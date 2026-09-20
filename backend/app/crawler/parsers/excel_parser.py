from pathlib import Path

from app.crawler.parsers.pdf_parser import ParsedDocument


class ExcelParser:
    def parse(self, path: Path) -> ParsedDocument:
        try:
            import openpyxl
        except ImportError:
            return ParsedDocument(text="", parser_type="openpyxl", error="Missing XLSX parser dependency: openpyxl")

        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"# {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        lines.append("\t".join(values))
            return ParsedDocument(text="\n".join(lines), parser_type="openpyxl")
        except Exception as exc:
            return ParsedDocument(text="", parser_type="openpyxl", error=str(exc))
