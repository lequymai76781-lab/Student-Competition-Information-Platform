from pathlib import Path
from zipfile import BadZipFile, ZipFile
import xml.etree.ElementTree as ET

from app.crawler.parsers.pdf_parser import ParsedDocument


class DocxParser:
    def parse(self, path: Path) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError:
            return self._parse_ooxml(path)

        try:
            document = Document(path)
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append("\t".join(cell.text.strip() for cell in row.cells))
            return ParsedDocument(text="\n".join(parts), parser_type="python-docx")
        except Exception:
            return self._parse_ooxml(path)

    def _parse_ooxml(self, path: Path) -> ParsedDocument:
        try:
            with ZipFile(path) as archive:
                xml = archive.read("word/document.xml")
            root = ET.fromstring(xml)
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            lines: list[str] = []
            for paragraph in root.iter(f"{namespace}p"):
                values = [node.text or "" for node in paragraph.iter(f"{namespace}t")]
                line = "".join(values).strip()
                if line:
                    lines.append(line)
            return ParsedDocument(text="\n".join(lines), parser_type="ooxml-stdlib")
        except (BadZipFile, KeyError, OSError, ET.ParseError) as exc:
            return ParsedDocument(text="", parser_type="ooxml-stdlib", error=str(exc))